"""One-time provenance importer for the locked Step 7 v1.1 endpoint catalog.

The v1.2 workbook branch is not treated as authority. Its typed schema document is
used only as a schema source, filtered to the exact 171 operations in the locked
v1.1 catalog. Run with explicit source paths; CI consumes only committed output.
"""

from __future__ import annotations

import argparse
from pathlib import Path
from typing import Any

import yaml
from openpyxl import load_workbook


HTTP_METHODS = {"get", "post", "put", "patch", "delete"}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--catalog", required=True, type=Path)
    parser.add_argument("--typed-source", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()

    catalog = read_catalog(args.catalog)
    with args.typed_source.open(encoding="utf-8-sig") as handle:
        specification: dict[str, Any] = yaml.safe_load(handle)

    authoritative_keys = {(row["method"], row["path"]) for row in catalog}
    for path in list(specification["paths"]):
        path_item = specification["paths"][path]
        for method in list(path_item):
            if method in HTTP_METHODS and (method, path) not in authoritative_keys:
                del path_item[method]
        if not any(method in HTTP_METHODS for method in path_item):
            del specification["paths"][path]

    operation_count = 0
    for row in catalog:
        operation = specification["paths"][row["path"]][row["method"]]
        operation["x-abp-module"] = row["module"]
        operation["x-abp-permission"] = row["permission"]
        operation["x-abp-scope-rule"] = row["scope_rule"]
        operation["x-abp-concurrency"] = row["concurrency"]
        operation["x-abp-idempotency"] = row["idempotency"]
        operation["x-abp-request-schema"] = row["request_schema"]
        operation["x-abp-success-schema"] = row["response_schema"]
        operation["x-abp-job-dto"] = "Job" in row["response_schema"]
        operation["x-abp-execution-mode"] = (
            "ASYNC" if operation["x-abp-job-dto"] else "SYNC"
        )
        operation["x-implementation-status"] = (
            "IMPLEMENTED"
            if row["module"] == "Platform"
            else "CONTRACTED_NOT_IMPLEMENTED"
        )
        operation["x-runtime-enabled"] = row["module"] == "Platform"
        operation["description"] = (
            f"Permission: {row['permission']}. Scope: {row['scope_rule']}. "
            f"{row['critical_note']}"
        )
        ensure_error_responses(operation)
        ensure_concurrency_contract(operation, row["concurrency"])
        if is_enterprise_list(row):
            apply_enterprise_list_query(operation, row["module"] == "Audit")
        operation_count += 1

    if operation_count != 171:
        raise RuntimeError(f"Expected 171 locked operations, found {operation_count}")

    specification["info"] = {
        "title": "ABP Commercial System API",
        "version": "1.1.0",
        "description": (
            "Authoritative OpenAPI 3.1 contract derived from the 171-operation "
            "Step 7 LOCKED v1.1 catalog. Contracted domain operations are not "
            "runtime-enabled until their implementation slice is delivered."
        ),
        "license": {
            "name": "Proprietary — internal ABP use only",
            "identifier": "LicenseRef-ABP-Proprietary",
        },
    }
    specification["jsonSchemaDialect"] = (
        "https://json-schema.org/draft/2020-12/schema"
    )
    for tag in specification.get("tags", []):
        tag["description"] = f"{tag['name']} V1 contract operations."
    specification["components"]["parameters"].update(common_parameters())
    specification["components"]["headers"] = common_headers()
    specification["components"]["schemas"].update(common_schemas())
    specification["components"]["responses"].update(common_responses())
    constrain_dynamic_json(specification)
    normalize_row_versions(specification)
    apply_locked_schema_guards(specification)
    apply_analytics_metadata(specification, catalog)
    prune_unused_components(specification, set(common_schemas()))

    args.output.parent.mkdir(parents=True, exist_ok=True)
    with args.output.open("w", encoding="utf-8", newline="\n") as handle:
        yaml.safe_dump(
            specification,
            handle,
            allow_unicode=True,
            sort_keys=False,
            width=100,
        )


def read_catalog(path: Path) -> list[dict[str, str]]:
    workbook = load_workbook(path, read_only=False, data_only=True)
    worksheet = workbook["11_Final_API_Catalog"]
    rows: list[dict[str, str]] = []
    for values in worksheet.iter_rows(min_row=2, values_only=True):
        if not values[0]:
            continue
        raw_path = text(values[1])
        rows.append(
            {
                "method": text(values[0]).lower(),
                "path": raw_path.removeprefix("/api/v1"),
                "module": text(values[2]),
                "purpose": text(values[3]),
                "permission": text(values[4]),
                "scope_rule": text(values[5]),
                "concurrency": text(values[6], "NONE"),
                "idempotency": text(values[7], "NONE"),
                "request_schema": text(values[8], "NONE"),
                "response_schema": text(values[9]),
                "critical_note": text(values[10]),
            }
        )
    return rows


def text(value: object, fallback: str = "") -> str:
    if value is None or str(value).strip() in {"", "—", "None"}:
        return fallback
    return str(value).strip()


def ensure_error_responses(operation: dict[str, Any]) -> None:
    responses = operation.setdefault("responses", {})
    for status in ["400", "401", "403", "404", "409", "412", "422", "429", "500"]:
        responses[status] = {"$ref": f"#/components/responses/Error{status}"}


def ensure_concurrency_contract(operation: dict[str, Any], concurrency: str) -> None:
    if concurrency == "If-Match":
        parameters = operation.setdefault("parameters", [])
        if not any(
            parameter.get("$ref") == "#/components/parameters/IfMatch"
            for parameter in parameters
            if isinstance(parameter, dict)
        ):
            parameters.append({"$ref": "#/components/parameters/IfMatch"})
    if concurrency in {"If-Match", "expectedRowVersion"}:
        for status, response in operation.get("responses", {}).items():
            if status.startswith("2") and isinstance(response, dict):
                response.setdefault("headers", {})["ETag"] = {
                    "$ref": "#/components/headers/ETag"
                }


def is_enterprise_list(row: dict[str, str]) -> bool:
    if row["method"] != "get":
        return False
    return any(
        marker in f"{row['response_schema']} {row['purpose']}".lower()
        for marker in ["list", "grid", "search", "timeline", "feed", "agenda"]
    )


def apply_enterprise_list_query(operation: dict[str, Any], cursor_feed: bool) -> None:
    existing = operation.setdefault("parameters", [])
    existing[:] = [
        parameter
        for parameter in existing
        if not (
            isinstance(parameter, dict)
            and parameter.get("$ref")
            in {
                "#/components/parameters/Cursor",
                "#/components/parameters/Page",
                "#/components/parameters/Offset",
                "#/components/parameters/Limit",
                "#/components/parameters/Sort",
                "#/components/parameters/Search",
            }
        )
    ]
    names = ["Cursor", "Limit"] if cursor_feed else ["Page", "Offset", "Limit", "Sort", "Search"]
    existing.extend({"$ref": f"#/components/parameters/{name}"} for name in names)
    operation["x-abp-pagination"] = "CURSOR" if cursor_feed else "OFFSET"


def common_parameters() -> dict[str, Any]:
    return {
        "Page": {
            "name": "page",
            "in": "query",
            "required": False,
            "schema": {"type": "integer", "minimum": 1, "default": 1},
        },
        "Offset": {
            "name": "offset",
            "in": "query",
            "required": False,
            "description": "Zero-based offset; when supplied it takes precedence over page.",
            "schema": {"type": "integer", "minimum": 0},
        },
        "Limit": {
            "name": "limit",
            "in": "query",
            "required": False,
            "schema": {"type": "integer", "minimum": 1, "maximum": 200, "default": 50},
        },
        "Sort": {
            "name": "sort",
            "in": "query",
            "required": False,
            "description": "Stable comma-separated field[:asc|desc] ordering; id is the final tie-breaker.",
            "schema": {
                "type": "string",
                "maxLength": 300,
                "pattern": r"^[A-Za-z][A-Za-z0-9]*(?::(?:asc|desc))?(?:,[A-Za-z][A-Za-z0-9]*(?::(?:asc|desc))?)*$",
            },
        },
        "Search": {
            "name": "search",
            "in": "query",
            "required": False,
            "schema": {"type": "string", "minLength": 1, "maxLength": 200},
        },
        "Cursor": {
            "name": "cursor",
            "in": "query",
            "required": False,
            "schema": {"type": "string", "minLength": 1, "maxLength": 1000},
        },
        "IfMatch": {
            "name": "If-Match",
            "in": "header",
            "required": True,
            "description": "Quoted ETag returned by the resource representation.",
            "schema": {"$ref": "#/components/schemas/RowVersion"},
        },
    }


def common_headers() -> dict[str, Any]:
    return {
        "ETag": {
            "description": "Opaque optimistic-concurrency version for the returned representation.",
            "schema": {"$ref": "#/components/schemas/RowVersion"},
        },
        "RetryAfter": {
            "description": "Seconds before a rate-limited request should be retried.",
            "schema": {"type": "integer", "minimum": 1},
        },
    }


def common_schemas() -> dict[str, Any]:
    scalar_param = {
        "oneOf": [
            {"type": "string", "maxLength": 500},
            {"type": "number"},
            {"type": "integer"},
            {"type": "boolean"},
            {"type": "null"},
        ]
    }
    return {
        "UUID": {"type": "string", "format": "uuid"},
        "Date": {"type": "string", "format": "date"},
        "DateTime": {"type": "string", "format": "date-time"},
        "CurrencyCode": {"type": "string", "pattern": r"^[A-Z]{3}$", "examples": ["USD", "EGP"]},
        "Decimal": {"type": "string", "pattern": r"^-?(?:0|[1-9]\d*)(?:\.\d{1,6})?$"},
        "Money": {
            "type": "object",
            "additionalProperties": False,
            "required": ["amount", "currency"],
            "properties": {
                "amount": {"$ref": "#/components/schemas/Decimal"},
                "currency": {"$ref": "#/components/schemas/CurrencyCode"},
            },
            "description": "Authoritative decimal amount transported as a string plus ISO currency.",
        },
        "Percentage": {"type": "string", "pattern": r"^(?:100(?:\.0{1,4})?|(?:\d{1,2})(?:\.\d{1,4})?)$"},
        "CountryCode": {"type": "string", "pattern": r"^[A-Z]{2}$"},
        "TraceId": {"type": "string", "minLength": 16, "maxLength": 128},
        "CorrelationId": {"type": "string", "minLength": 8, "maxLength": 128},
        "RowVersion": {
            "type": "string",
            "minLength": 3,
            "maxLength": 200,
            "pattern": '^W?/\"[^\"]+\"$',
            "description": "Opaque quoted ETag. Clients must not parse its contents.",
        },
        "ResourceReference": {
            "type": "object",
            "additionalProperties": False,
            "required": ["resourceType", "resourceId"],
            "properties": {
                "resourceType": {"type": "string", "minLength": 1, "maxLength": 100},
                "resourceId": {"type": "string", "minLength": 1, "maxLength": 200},
            },
        },
        "DynamicJsonValue": {
            "oneOf": [
                {"type": "string", "maxLength": 4000},
                {"type": "number"},
                {"type": "boolean"},
                {"type": "null"},
                {"type": "array", "maxItems": 100, "items": {"$ref": "#/components/schemas/DynamicJsonValue"}},
                {
                    "type": "object",
                    "maxProperties": 100,
                    "propertyNames": {"pattern": r"^[A-Za-z][A-Za-z0-9_.-]{0,99}$"},
                    "additionalProperties": {"$ref": "#/components/schemas/DynamicJsonValue"},
                },
            ],
            "description": "Bounded schemaless source/snapshot metadata; never an authoritative business DTO.",
        },
        "ErrorParamMap": {
            "type": "object",
            "maxProperties": 20,
            "propertyNames": {"pattern": r"^[A-Za-z][A-Za-z0-9_.-]{0,63}$"},
            "additionalProperties": scalar_param,
        },
        "FieldError": {
            "type": "object",
            "additionalProperties": False,
            "required": ["field", "code", "messageKey", "params"],
            "properties": {
                "field": {"type": "string", "minLength": 1, "maxLength": 200},
                "code": {"type": "string", "minLength": 1, "maxLength": 100},
                "messageKey": {"type": "string", "minLength": 1, "maxLength": 200},
                "params": {"$ref": "#/components/schemas/ErrorParamMap"},
            },
        },
        "RuleError": {
            "type": "object",
            "additionalProperties": False,
            "required": ["code", "messageKey", "params"],
            "properties": {
                "code": {"type": "string", "minLength": 1, "maxLength": 100},
                "messageKey": {"type": "string", "minLength": 1, "maxLength": 200},
                "params": {"$ref": "#/components/schemas/ErrorParamMap"},
            },
        },
        "ConflictMetadata": {
            "type": ["object", "null"],
            "additionalProperties": False,
            "required": ["resource", "expectedRowVersion", "currentRowVersion"],
            "properties": {
                "resource": {"$ref": "#/components/schemas/ResourceReference"},
                "expectedRowVersion": {"$ref": "#/components/schemas/RowVersion"},
                "currentRowVersion": {"$ref": "#/components/schemas/RowVersion"},
            },
        },
        "ErrorBody": {
            "type": "object",
            "additionalProperties": False,
            "required": ["code", "messageKey", "params", "fieldErrors", "ruleErrors", "traceId", "correlationId", "conflict"],
            "properties": {
                "code": {"type": "string", "minLength": 1, "maxLength": 100},
                "messageKey": {"type": "string", "minLength": 1, "maxLength": 200},
                "params": {"$ref": "#/components/schemas/ErrorParamMap"},
                "fieldErrors": {"type": "array", "maxItems": 100, "items": {"$ref": "#/components/schemas/FieldError"}},
                "ruleErrors": {"type": "array", "maxItems": 100, "items": {"$ref": "#/components/schemas/RuleError"}},
                "traceId": {"$ref": "#/components/schemas/TraceId"},
                "correlationId": {"$ref": "#/components/schemas/CorrelationId"},
                "conflict": {"$ref": "#/components/schemas/ConflictMetadata"},
            },
        },
        "ErrorEnvelope": {
            "type": "object",
            "additionalProperties": False,
            "required": ["error"],
            "properties": {"error": {"$ref": "#/components/schemas/ErrorBody"}},
        },
        "OffsetPaginationMeta": {
            "type": "object",
            "additionalProperties": False,
            "required": ["page", "offset", "limit", "total", "hasNext"],
            "properties": {
                "page": {"type": "integer", "minimum": 1},
                "offset": {"type": "integer", "minimum": 0},
                "limit": {"type": "integer", "minimum": 1, "maximum": 200},
                "total": {"type": ["integer", "null"], "minimum": 0},
                "hasNext": {"type": "boolean"},
            },
        },
        "CursorPaginationMeta": {
            "type": "object",
            "additionalProperties": False,
            "required": ["nextCursor", "hasMore", "limit"],
            "properties": {
                "nextCursor": {"type": ["string", "null"], "maxLength": 1000},
                "hasMore": {"type": "boolean"},
                "limit": {"type": "integer", "minimum": 1, "maximum": 200},
            },
        },
        "PaginationMeta": {"$ref": "#/components/schemas/OffsetPaginationMeta"},
        "PermissionCode": {"type": "string", "pattern": r"^[a-z][a-z0-9.-]{2,99}$"},
        "RoleSummary": {
            "type": "object",
            "additionalProperties": False,
            "required": ["roleId", "code", "name", "workspace"],
            "properties": {
                "roleId": {"$ref": "#/components/schemas/UUID"},
                "code": {"type": "string", "minLength": 1, "maxLength": 100},
                "name": {"type": "string", "minLength": 1, "maxLength": 200},
                "workspace": {"$ref": "#/components/schemas/WorkspaceCode"},
            },
        },
        "WorkspaceCode": {"type": "string", "enum": ["SALES", "LOGISTICS", "OPERATIONS", "FINANCE", "ADMIN"]},
        "WorkspaceAccess": {
            "type": "object",
            "additionalProperties": False,
            "required": ["workspace", "enabled", "capabilities"],
            "properties": {
                "workspace": {"$ref": "#/components/schemas/WorkspaceCode"},
                "enabled": {"type": "boolean"},
                "capabilities": {"type": "array", "uniqueItems": True, "items": {"$ref": "#/components/schemas/PermissionCode"}},
            },
        },
        "CountryScope": {
            "type": "object",
            "additionalProperties": False,
            "required": ["enforced", "countryCodes"],
            "properties": {
                "enforced": {"type": "boolean", "const": True},
                "countryCodes": {"type": "array", "uniqueItems": True, "items": {"$ref": "#/components/schemas/CountryCode"}},
            },
        },
        "SessionSummary": {
            "type": "object",
            "additionalProperties": False,
            "required": ["sessionId", "authTransport", "assurance", "issuedAt", "expiresAt"],
            "properties": {
                "sessionId": {"type": "string", "minLength": 8, "maxLength": 200},
                "authTransport": {"type": "string", "enum": ["BFF_COOKIE", "DESKTOP_OAUTH_PKCE"]},
                "assurance": {"type": "string", "enum": ["PASSWORD", "MFA", "FEDERATED_MFA"]},
                "issuedAt": {"$ref": "#/components/schemas/DateTime"},
                "expiresAt": {"$ref": "#/components/schemas/DateTime"},
            },
        },
        "StepUpRequirementState": {
            "type": "object",
            "additionalProperties": False,
            "required": ["required", "satisfied", "permittedMethods"],
            "properties": {
                "required": {"type": "boolean"},
                "satisfied": {"type": "boolean"},
                "permittedMethods": {"type": "array", "items": {"type": "string", "enum": ["MFA", "REAUTHENTICATE"]}},
            },
        },
        "EffectiveAccessContext": {
            "type": "object",
            "additionalProperties": False,
            "required": ["roles", "permissions", "countryScope", "workspaces", "policyVersion"],
            "properties": {
                "roles": {"type": "array", "items": {"$ref": "#/components/schemas/RoleSummary"}},
                "permissions": {"type": "array", "uniqueItems": True, "items": {"$ref": "#/components/schemas/PermissionCode"}},
                "countryScope": {"$ref": "#/components/schemas/CountryScope"},
                "workspaces": {"type": "array", "items": {"$ref": "#/components/schemas/WorkspaceAccess"}},
                "policyVersion": {"type": "integer", "minimum": 1},
            },
        },
        "Me": {
            "type": "object",
            "additionalProperties": False,
            "required": ["user", "session", "access", "stepUp"],
            "properties": {
                "user": {"$ref": "#/components/schemas/User"},
                "session": {"$ref": "#/components/schemas/SessionSummary"},
                "access": {"$ref": "#/components/schemas/EffectiveAccessContext"},
                "stepUp": {"$ref": "#/components/schemas/StepUpRequirementState"},
            },
            "description": "Sales Management and Sales Support are roles/capabilities within SALES, never separate workspaces.",
        },
        "JobStatus": {"type": "string", "enum": ["PENDING", "PROCESSING", "SUCCEEDED", "FAILED", "DEAD"]},
        "JobError": {
            "type": ["object", "null"],
            "additionalProperties": False,
            "required": ["code", "messageKey", "traceId"],
            "properties": {
                "code": {"type": "string", "minLength": 1, "maxLength": 100},
                "messageKey": {"type": "string", "minLength": 1, "maxLength": 200},
                "traceId": {"$ref": "#/components/schemas/TraceId"},
            },
        },
        "Job": {
            "type": "object",
            "additionalProperties": False,
            "required": ["jobId", "jobType", "status", "progress", "attemptCount", "maxAttempts", "createdAt", "startedAt", "finishedAt", "resultReference", "error", "traceId", "correlationId"],
            "properties": {
                "jobId": {"$ref": "#/components/schemas/UUID"},
                "jobType": {"type": "string", "minLength": 1, "maxLength": 100},
                "status": {"$ref": "#/components/schemas/JobStatus"},
                "progress": {"type": "integer", "minimum": 0, "maximum": 100},
                "attemptCount": {"type": "integer", "minimum": 0},
                "maxAttempts": {"type": "integer", "minimum": 1},
                "createdAt": {"$ref": "#/components/schemas/DateTime"},
                "startedAt": {"oneOf": [{"$ref": "#/components/schemas/DateTime"}, {"type": "null"}]},
                "finishedAt": {"oneOf": [{"$ref": "#/components/schemas/DateTime"}, {"type": "null"}]},
                "resultReference": {"oneOf": [{"$ref": "#/components/schemas/ResourceReference"}, {"type": "null"}]},
                "error": {"$ref": "#/components/schemas/JobError"},
                "traceId": {"$ref": "#/components/schemas/TraceId"},
                "correlationId": {"$ref": "#/components/schemas/CorrelationId"},
            },
        },
        "AnalyticsResponseMetadata": {
            "type": "object",
            "additionalProperties": False,
            "required": ["metricId", "metricVersion", "asOf", "source", "currency", "basis", "scope"],
            "properties": {
                "metricId": {"type": "string", "minLength": 1, "maxLength": 100},
                "metricVersion": {"type": "string", "minLength": 1, "maxLength": 50},
                "asOf": {"$ref": "#/components/schemas/DateTime"},
                "source": {"type": "string", "minLength": 1, "maxLength": 100},
                "currency": {"oneOf": [{"$ref": "#/components/schemas/CurrencyCode"}, {"type": "null"}]},
                "basis": {"type": "string", "minLength": 1, "maxLength": 200},
                "scope": {"type": "string", "minLength": 1, "maxLength": 200},
            },
            "description": "Semantic provenance only; the frontend is never the KPI authority.",
        },
        "Health": {
            "type": "object",
            "additionalProperties": False,
            "required": ["status", "timestamp", "checks"],
            "properties": {
                "status": {"type": "string", "enum": ["UP", "DOWN", "DEGRADED"]},
                "timestamp": {"$ref": "#/components/schemas/DateTime"},
                "checks": {
                    "type": "array",
                    "maxItems": 20,
                    "items": {
                        "type": "object",
                        "additionalProperties": False,
                        "required": ["name", "status"],
                        "properties": {
                            "name": {"type": "string", "minLength": 1, "maxLength": 100},
                            "status": {"type": "string", "enum": ["UP", "DOWN"]},
                        },
                    },
                },
            },
        },
    }


def common_responses() -> dict[str, Any]:
    descriptions = {
        "400": "Request validation failed",
        "401": "No valid authenticated session",
        "403": "Authenticated caller is forbidden",
        "404": "Resource not found or safely masked",
        "409": "Business-state conflict",
        "412": "Optimistic-concurrency precondition failed",
        "422": "Business rule violation",
        "429": "Rate limit exceeded",
        "500": "Safe internal error",
        "503": "Dependency temporarily unavailable",
    }
    result: dict[str, Any] = {}
    for status, description in descriptions.items():
        response: dict[str, Any] = {
            "description": description,
            "content": {
                "application/json": {
                    "schema": {"$ref": "#/components/schemas/ErrorEnvelope"}
                }
            },
        }
        if status == "429":
            response["headers"] = {
                "Retry-After": {"$ref": "#/components/headers/RetryAfter"}
            }
        result[f"Error{status}"] = response
    return result


def constrain_dynamic_json(node: Any) -> None:
    if isinstance(node, dict):
        for key, value in list(node.items()):
            if key == "additionalProperties" and value is True:
                node[key] = {"$ref": "#/components/schemas/DynamicJsonValue"}
                node.setdefault("maxProperties", 100)
            else:
                constrain_dynamic_json(value)
    elif isinstance(node, list):
        for value in node:
            constrain_dynamic_json(value)


def normalize_row_versions(node: Any) -> None:
    if isinstance(node, dict):
        properties = node.get("properties")
        if isinstance(properties, dict):
            for name in ["rowVersion", "expectedRowVersion"]:
                if name in properties:
                    description = properties[name].get("description") if isinstance(properties[name], dict) else None
                    properties[name] = {"$ref": "#/components/schemas/RowVersion"}
                    if description:
                        properties[name]["description"] = description
        for value in node.values():
            normalize_row_versions(value)
    elif isinstance(node, list):
        for value in node:
            normalize_row_versions(value)


def apply_locked_schema_guards(specification: dict[str, Any]) -> None:
    schemas = specification["components"]["schemas"]
    capacity_rate = schemas.get("CapacityRate")
    if capacity_rate:
        capacity_rate["properties"]["categoryCode"] = {
            "type": "string",
            "enum": ["STANDARD", "ORDINARY", "CURVED"],
        }
        capacity_rate["allOf"] = [
            {
                "if": {"properties": {"categoryCode": {"const": category}}},
                "then": {"properties": {"boardsPerDay": {"const": rate}}},
            }
            for category, rate in [
                ("STANDARD", "900"),
                ("ORDINARY", "700"),
                ("CURVED", "600"),
            ]
        ]
        capacity_rate["description"] = (
            "Estimate-only rates at 24 hours/day. No reservation, overbooking, work order, MES, routing, or QC semantics."
        )

    returned_scrap = schemas.get("ReturnedScrapCase")
    if returned_scrap:
        returned_scrap["properties"]["appliedCreditPctSnapshot"] = {
            "type": "string",
            "const": "25",
            "description": "Locked 25 percent benefit rate snapshot.",
        }
        returned_scrap["properties"]["benefitTypeCode"] = {
            "type": "string",
            "enum": ["CREDIT", "REPLACEMENT"],
        }
        returned_scrap["description"] = (
            "Manager-approved benefit = 25% × original purchase unit price × eligible returned quantity. Current product price and Net Sales are excluded."
        )

    quotation = schemas.get("Quotation")
    if quotation:
        quotation["description"] = (
            "No formal revisions: DRAFT is mutable; SENT is immutable; a material post-SENT change creates a new quotation linked by supersedesQuotationId. Approval binds to the exact materialContentHash snapshot."
        )

    source_order = schemas.get("SourceSalesOrder")
    if source_order:
        source_order["description"] = (
            "Imported read-only ERP canonical fact with source lineage. No internal create, confirm, amend, handover, close, shipment, or POD command semantics."
        )

    logistics_version = schemas.get("LogisticsPlanVersion")
    if logistics_version:
        logistics_version["description"] = (
            "Planning-only Flat-on-Face scenario: footprint is length × width, thickness is stack height, and face rotation is 0/90. No shipment, dispatch, delivery, or POD execution."
        )


def apply_analytics_metadata(specification: dict[str, Any], catalog: list[dict[str, str]]) -> None:
    schemas = specification["components"]["schemas"]
    analytics_modules = {"Analytics", "Dashboard", "Targets", "AI"}
    for row in catalog:
        if row["module"] not in analytics_modules:
            continue
        operation = specification["paths"][row["path"]][row["method"]]
        for status, response in operation.get("responses", {}).items():
            if not status.startswith("2") or not isinstance(response, dict):
                continue
            reference = (
                response.get("content", {})
                .get("application/json", {})
                .get("schema", {})
                .get("$ref")
            )
            if not isinstance(reference, str) or not reference.startswith("#/components/schemas/"):
                continue
            schema_name = reference.rsplit("/", 1)[-1]
            envelope = schemas.get(schema_name)
            if isinstance(envelope, dict) and isinstance(envelope.get("properties"), dict):
                envelope["properties"]["analyticsMetadata"] = {
                    "$ref": "#/components/schemas/AnalyticsResponseMetadata"
                }
                required = envelope.setdefault("required", [])
                if "analyticsMetadata" not in required:
                    required.append("analyticsMetadata")


def prune_unused_components(specification: dict[str, Any], keep_schemas: set[str]) -> None:
    components = specification["components"]
    used: set[tuple[str, str]] = set()

    def collect_references(node: Any) -> None:
        if isinstance(node, dict):
            reference = node.get("$ref")
            if isinstance(reference, str) and reference.startswith("#/components/"):
                parts = reference.split("/")
                if len(parts) >= 4:
                    used.add((parts[2], parts[3]))
            for value in node.values():
                collect_references(value)
        elif isinstance(node, list):
            for value in node:
                collect_references(value)

    collect_references(specification["paths"])
    for schema_name in keep_schemas:
        used.add(("schemas", schema_name))

    processed: set[tuple[str, str]] = set()
    while processed != used:
        pending = used - processed
        for category, name in pending:
            component = components.get(category, {}).get(name)
            if component is not None:
                collect_references(component)
        processed.update(pending)

    for category, entries in list(components.items()):
        if category == "securitySchemes" or not isinstance(entries, dict):
            continue
        for name in list(entries):
            if (category, name) not in used:
                del entries[name]


if __name__ == "__main__":
    main()
